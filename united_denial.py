import os
import re
import sys
import struct, time
import subprocess
from datetime import date
import datetime
import openpyxl
import pdftotext
import time
import requests

from make_log import log_exceptions
from custom_parallel import write
from custom_datadict import make_datadict


'''
wbkName = 'log file.xlsx'
wbk= openpyxl.load_workbook(wbkName)
s1=wbk.worksheets[0]
s2=wbk.worksheets[1]
row_count_1 = s2.max_row
s2.cell(row_count_1+1, column=1).value=sys.argv[2]
s2.cell(row_count_1+1, column=2).value=sys.argv[3]
s2.cell(row_count_1+1, column=3).value=sys.argv[4]
s2.cell(row_count_1+1, column=5).value=now
s2.cell(row_count_1+1, column=7).value=sys.argv[5]
s2.cell(row_count_1+1, column=8).value=sys.argv[6]
'''

subprocess.run(["python", "updation.py","1","max1","1",sys.argv[2]])
subprocess.run(["python", "updation.py","1","max","2",sys.argv[3]])
subprocess.run(["python", "updation.py","1","max","3",sys.argv[4]])
subprocess.run(["python", "updation.py","1","max","5",str(now)])
subprocess.run(["python", "updation.py","1","max","7",sys.argv[5]])
subprocess.run(["python", "updation.py","1","max","8",sys.argv[6]])

with open(sys.argv[1], "rb") as f:
    pdf = pdftotext.PDF(f)

with open('united/output1.txt', 'w') as f:
    f.write(" ".join(pdf))     
with open('united/output1.txt', 'r') as myfile:
    f = myfile.read()

# with open('nurseline/output.txt', 'w') as f:
#   f.write(" ".join(pdf))     
# with open('nurseline/output.txt', 'r') as myfile:
#   f = myfile.read()


try:
    datadict = make_datadict(f)
    data = [i for i in sys.argv[1:]]
    data2 = [datadict[i] for i in datadict]
    data.extend(data2)
    data3 = str(datadict).replace('{', '\{').replace('}', '\}')
    data.append(data3)
    write(data)
except:
    log_exceptions()